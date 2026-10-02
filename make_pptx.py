#!/usr/bin/env python3
"""Generate a PowerPoint presentation about SmartVAPT."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ---- Palette (dark cyber theme) ----
BG_DARK   = RGBColor(0x12, 0x1A, 0x27)   # deep navy
ACCENT    = RGBColor(0x00, 0xB4, 0xD8)   # cyan
ACCENT2   = RGBColor(0x2E, 0xCC, 0x71)   # green
WARN      = RGBColor(0xFF, 0xB0, 0x20)   # amber
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
GREY      = RGBColor(0xB8, 0xC2, 0xCE)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(title_text=None, subtitle_text=None):
    slide = prs.slides.add_slide(BLANK)
    # background
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BG_DARK; bg.line.fill.background()
    slide.shapes._spTree.remove(bg._element); slide.shapes._spTree.insert(2, bg._element)
    # accent bar
    bar = slide.shapes.add_shape(1, Inches(0.5), Inches(0.45), Inches(0.12), Inches(0.9))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    if title_text:
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.8), Inches(0.8))
        p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = title_text
        r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = WHITE
        r.font.name = "Calibri"
    if subtitle_text:
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(11.8), Inches(0.5))
        p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = subtitle_text
        r.font.size = Pt(16); r.font.color.rgb = GREY; r.font.italic = True
    return slide


def bullets(slide, items, left=0.9, top=1.7, width=11.6, height=5.3, size=18):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for text, level in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(8)
        r = p.add_run(); r.text = text
        r.font.size = Pt(size - level * 2)
        r.font.color.rgb = WHITE if level == 0 else GREY
        r.font.name = "Calibri"
        if level == 0 and text.endswith(":"):
            r.font.color.rgb = ACCENT; r.font.bold = True
    return tb


# ================= SLIDE 1 — TITLE =================
s = add_slide()
tb = s.shapes.add_textbox(Inches(1), Inches(2.3), Inches(11.3), Inches(1.6))
p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = "SmartVAPT"
r.font.size = Pt(66); r.font.bold = True; r.font.color.rgb = ACCENT
tb2 = s.shapes.add_textbox(Inches(1), Inches(3.9), Inches(11.3), Inches(1.2))
p = tb2.text_frame.paragraphs[0]; r = p.add_run()
r.text = "Intelligent Web Application Vulnerability Assessment & Penetration Testing"
r.font.size = Pt(24); r.font.color.rgb = WHITE
p2 = tb2.text_frame.add_paragraph(); r2 = p2.add_run()
r2.text = "Scan \u2192 Discover \u2192 Exploit-aware \u2192 Auto-Remediate \u2192 Report"
r2.font.size = Pt(16); r2.font.color.rgb = ACCENT2
bar = s.shapes.add_shape(1, Inches(1), Inches(3.75), Inches(4.5), Inches(0.05))
bar.fill.solid(); bar.fill.fore_color.rgb = WARN; bar.line.fill.background()

# ================= SLIDE 2 — WHAT IS SMARTVAPT =================
s = add_slide("What is SmartVAPT?", "An open-source, Python-based security automation platform")
bullets(s, [
    ("SmartVAPT is an end-to-end VAPT engine that automates the full security-testing lifecycle for web applications:", 0),
    ("Discovery: crawls targets, fingerprints tech stacks, and enumerates endpoints", 1),
    ("Assessment: runs plugin-based vulnerability checks (SQLi, XSS, CSRF, SSRF, IDOR, headers, TLS, CVEs...)", 1),
    ("Validation: confirms findings with safe proof-of-concept evidence before reporting", 1),
    ("Remediation: generates AND APPLIES code fixes for detected vulnerabilities", 1),
    ("Reporting: produces professional PDF/HTML reports with CVSS scoring", 1),
    ("Unlike traditional scanners that only report problems, SmartVAPT closes the loop by patching them.", 0),
])

# ================= SLIDE 3 — HOW IT WORKS =================
s = add_slide("How It Works", "The automated scan-to-fix pipeline")
stages = [
    ("1. TARGET", "Input URL / repo;\nauth & scope setup"),
    ("2. CRAWL", "Endpoint discovery,\ntech fingerprinting"),
    ("3. SCAN", "Plugin engine runs\nCVE & OWASP checks"),
    ("4. VALIDATE", "PoC verification,\nfalse-positive trim"),
    ("5. FIX", "Auto-generates &\napplies patches"),
    ("6. REPORT", "PDF/HTML output,\nCVSS + retest"),
]
x = 0.55
for name, desc in stages:
    box = s.shapes.add_shape(5, Inches(x), Inches(2.2), Inches(1.85), Inches(2.2))  # rounded rect
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x1C, 0x2A, 0x3D)
    box.line.color.rgb = ACCENT; box.line.width = Pt(1.5)
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = name
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = desc
    r2.font.size = Pt(11); r2.font.color.rgb = WHITE
    x += 2.1
tb = s.shapes.add_textbox(Inches(0.9), Inches(4.9), Inches(11.6), Inches(1.6))
tf = tb.text_frame; tf.word_wrap = True
p = tf.paragraphs[0]; r = p.add_run()
r.text = "Headless mode: run_full_scan.py executes the entire pipeline from the CLI — no GUI required."
r.font.size = Pt(15); r.font.color.rgb = GREY
p2 = tf.add_paragraph(); r2 = p2.add_run()
r2.text = "Regression check: patched CVEs are automatically removed from the final findings list."
r2.font.size = Pt(15); r2.font.color.rgb = ACCENT2

# ================= SLIDE 4 — KEY FEATURES =================
s = add_slide("Key Features", "Everything under one roof")
bullets(s, [
    ("Plugin Architecture:", 0),
    ("Modular checks — add new vulnerability tests without touching the core engine", 1),
    ("CVE Intelligence:", 0),
    ("Matches detected software versions against known CVEs with severity ratings", 1),
    ("Auto-Remediation Engine:", 0),
    ("Generates code-level fixes (e.g., parameterized queries, secure headers) and applies them safely", 1),
    ("Dual Interface:", 0),
    ("Desktop GUI for interactive testing + headless CLI for CI/CD pipelines", 1),
    ("Professional Reporting:", 0),
    ("PDF & HTML reports with executive summaries, technical details, and remediation guidance", 1),
    ("Extensibility:", 0),
    ("Python API — integrate scans into your own tooling and automation", 1),
], size=16)

# ================= SLIDE 5 — ADVANTAGES =================
s = add_slide("Advantages", "Why choose SmartVAPT?")
adv = [
    ("Scan + Fix in One Tool", "Most scanners stop at detection; SmartVAPT remediates too."),
    ("Open Source & Free", "No license fees, full transparency, community-driven."),
    ("CI/CD Ready", "Headless scans fit straight into build pipelines."),
    ("Low False Positives", "Findings are validated with proof-of-concept evidence."),
    ("Fast Time-to-Fix", "Instant patch suggestions cut remediation cycles to minutes."),
    ("Lightweight", "Pure Python stack — runs on a laptop, no heavy infrastructure."),
]
positions = [(0.7, 1.9), (4.85, 1.9), (9.0, 1.9), (0.7, 4.35), (4.85, 4.35), (9.0, 4.35)]
for (title, body), (px, py) in zip(adv, positions):
    card = s.shapes.add_shape(5, Inches(px), Inches(py), Inches(3.7), Inches(2.1))
    card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0x1C, 0x2A, 0x3D)
    card.line.color.rgb = ACCENT2; card.line.width = Pt(1.25)
    tf = card.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; r = p.add_run(); r.text = title
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = ACCENT2
    p2 = tf.add_paragraph(); r2 = p2.add_run(); r2.text = body
    r2.font.size = Pt(12); r2.font.color.rgb = WHITE

# ================= SLIDE 6 — USE CASES =================
s = add_slide("Use Cases", "Who uses SmartVAPT and how")
bullets(s, [
    ("Development Teams:", 0),
    ("Pre-commit and pre-deploy security checks inside CI/CD (GitHub Actions, Jenkins, GitLab CI)", 1),
    ("Security Professionals / Pentesters:", 0),
    ("Rapid initial assessments, then manual deep-dives using discovered endpoints", 1),
    ("Freelancers & Agencies:", 0),
    ("Deliver branded PDF audit reports to clients quickly", 1),
    ("Students & Educators:", 0),
    ("Hands-on lab for learning OWASP Top 10, CVE analysis, and secure coding", 1),
    ("Small & Medium Businesses:", 0),
    ("Affordable, automated security posture monitoring without enterprise scanner costs", 1),
    ("Compliance & Audits:", 0),
    ("Recurring scans and documented evidence for ISO 27001, SOC 2, PCI-DSS readiness", 1),
], size=16)

# ================= SLIDE 7 — TECH STACK =================
s = add_slide("Technical Overview", "Built with a modern Python security stack")
rows = [
    ("Language", "Python 3 (modular, object-oriented design)"),
    ("Scanning Engine", "Plugin-based checker system (OWASP-aligned)"),
    ("Crawling", "Automated endpoint discovery & tech fingerprinting"),
    ("Vuln Data", "CVE matching with CVSS severity scoring"),
    ("Remediation", "Patch generator + safe in-place code fixer"),
    ("Reporting", "PDF / HTML export with executive & technical sections"),
    ("Interfaces", "Desktop GUI + headless CLI for automation"),
    ("Testing", "pytest suite (38+ tests) with regression coverage"),
]
top = 1.8
for i, (k, v) in enumerate(rows):
    row = s.shapes.add_shape(1, Inches(0.9), Inches(top + i * 0.62), Inches(11.5), Inches(0.55))
    row.fill.solid()
    row.fill.fore_color.rgb = RGBColor(0x1C, 0x2A, 0x3D) if i % 2 == 0 else RGBColor(0x16, 0x20, 0x30)
    row.line.fill.background()
    tf = row.text_frame; tf.margin_left = Inches(0.15); tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = f"{k}:  "
    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = ACCENT
    r2 = p.add_run(); r2.text = v
    r2.font.size = Pt(14); r2.font.color.rgb = WHITE

# ================= SLIDE 8 — SUMMARY =================
s = add_slide("Summary", "SmartVAPT at a glance")
bullets(s, [
    ("What: an open-source, all-in-one VAPT platform for web applications.", 0),
    ("Differentiator: the only step most tools skip — automated remediation of found issues.", 0),
    ("Best for: developers, pentesters, agencies, educators, and SMBs needing affordable automation.", 0),
    ("Workflow: Crawl \u2192 Scan \u2192 Validate \u2192 Fix \u2192 Report \u2014 fully scriptable from the CLI.", 0),
    ("Next steps: try a headless demo scan, extend it with a custom plugin, or wire it into your CI pipeline.", 0),
], size=19)
tb = s.shapes.add_textbox(Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.8))
p = tb.text_frame.paragraphs[0]; r = p.add_run()
r.text = "Thank you!  Questions?"
r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = WARN

prs.save("/workspace/SmartVAPT_Presentation.pptx")
print("Saved SmartVAPT_Presentation.pptx with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
